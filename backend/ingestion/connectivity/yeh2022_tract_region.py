"""Conectividad estructural tracto->región derivada de Yeh (2022, Nature
Communications, "Population-based tract-to-region connectome of the
human brain and its hierarchical topology", DOI 10.1038/s41467-022-32595-4):
para cada uno de 26 tractos de sustancia blanca nombrados y cada uno de
los 180 hemisferios-área de la parcelación HCP-MMP1.0, la probabilidad
poblacional (0-1, sobre 1065 sujetos) de que ese tracto atraviese esa
región cortical.

Decisión (con la usuaria, 29/08/2026, ver docs/analisis-arquitectura.md):
NO se deriva una matriz región-región a partir de estos datos. El
archivo de origen no contiene una afirmación directa de "la región A se
conecta con la región B": contiene una afirmación de "el tracto T
atraviesa la región R con probabilidad P". Colapsar esto en conexiones
región-región (p. ej. por co-ocurrencia en un mismo tracto) sería un
dato INFERIDO por nosotras, no observado por el estudio. Se ha optado
por cargarlo tal cual: el tracto es una entidad `Tract` propia del grafo
(ya prevista en el esquema, sección 6), con conexiones tracto -> región.

Cada hemisferio es una entidad `Tract` propia (mismo criterio ya
aplicado a `Region` en hcp_mmp1.py: son estructuras físicamente
distintas, no una sola entidad con un atributo de lado).

Corrección de nomenclatura verificada empíricamente sobre los archivos
reales (no asumida): la matriz de origen usa las abreviaturas "PTAT" y
"C_R" en la cabecera de columnas, mientras que la tabla de abreviaturas
que acompaña al mismo paper (Supplementary Table 1, abbreviation2.xlsx)
usa "TPAT" (Temporo-Parietal Aslant Tract) y "C_PR" (Cingulum,
Parolfactory Segment). De los 26 códigos de cada archivo, 24 coinciden
literalmente; solo estos 2 difieren en cada lado, y por proceso de
eliminación (ningún otro código sobra ni falta) son el mismo tracto en
ambos casos. Se documenta la correspondencia explícitamente en
`_HEADER_TO_ABBREVIATION_TABLE_CODE` en vez de asumirla en silencio.

No se aplica ningún umbral al guardar los datos (mismo criterio que
Brainnetome, sección 24): se cargan las 9360 combinaciones tracto-región
reales (26 tractos x 2 hemisferios x 180 áreas), incluidas las que
tienen probabilidad 0.0 (el 75.3% de la matriz, verificado
empíricamente el 29/08/2026) -- es el filtro de peso mínimo ya
existente el que decide qué se muestra, nunca la carga de datos.

evidence_level='indirect' (mismo criterio que la conectividad
estructural de Brainnetome en brainnetome_sc.py): es tractografía
probabilística de difusión, no trazado de vías confirmado, aunque esté
agregada sobre 1065 sujetos reales.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from backend.ontology.schema import EntityType, build_id

METHOD = (
    "probabilidad_poblacional_de_interseccion_tracto_region_yeh_2022_"
    "sobre_1065_sujetos_sin_umbral"
)

EXPECTED_HEADER_TRACT_CODES: tuple[str, ...] = (
    "FAT", "AF", "PTAT", "MdLF", "SLF_III", "SLF_II", "SLF_I", "C_FP",
    "C_R", "C_FPH", "C_PHP", "C_PH", "ILF", "IFOF", "UF", "VOF", "CST",
    "CBT", "CStr_A", "CStr_P", "CStr_S", "CTh_A", "CTh_P", "CTh_S", "OR", "F",
)

_HEADER_TO_ABBREVIATION_TABLE_CODE: dict[str, str] = {
    "PTAT": "TPAT",
    "C_R": "C_PR",
}

_HEMISPHERE_LABEL = {"L": "izquierdo", "R": "derecho"}


@dataclass(frozen=True)
class TractDefinition:
    id: str
    name: str
    hemisphere: str  # "L" o "R"
    header_code: str  # código tal como aparece en la matriz de origen


@dataclass(frozen=True)
class TractRegionConnection:
    id: str
    source_id: str  # id del Tract
    target_id: str  # id de la Region (hcp-mmp1)
    weight: float
    hemisphere: str


def read_tract_full_names(abbreviation_xlsx_path: Path) -> dict[str, str]:
    """Lee la tabla de abreviaturas (Supplementary Table 1) y devuelve
    {código: nombre completo}. Salta las dos filas de cabecera reales
    del archivo ("Supplmentary Table 1" y la fila de nombres de
    columna)."""
    import openpyxl

    wb = openpyxl.load_workbook(str(abbreviation_xlsx_path), read_only=True, data_only=True)
    ws = wb["Sheet1"]
    names: dict[str, str] = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        code, full_name = row[0], row[1]
        if code is None:
            continue
        names[str(code).strip()] = str(full_name).strip()
    wb.close()
    return names


def tract_definitions(full_names: dict[str, str]) -> list[TractDefinition]:
    """Construye las 52 entidades `Tract` (26 tractos x 2 hemisferios).
    Lanza `ValueError` si algún código de cabecera no tiene, tras aplicar
    la corrección conocida, un nombre completo en la tabla de
    abreviaturas -- nunca se registra un tracto sin nombre real."""
    definitions: list[TractDefinition] = []
    for header_code in EXPECTED_HEADER_TRACT_CODES:
        lookup_code = _HEADER_TO_ABBREVIATION_TABLE_CODE.get(header_code, header_code)
        if lookup_code not in full_names:
            raise ValueError(
                f"código de tracto {header_code!r} (buscado como "
                f"{lookup_code!r}) no está en la tabla de abreviaturas"
            )
        full_name = full_names[lookup_code]
        for hemisphere in ("L", "R"):
            local_code = f"{hemisphere}_{header_code}".lower()
            tract_id = build_id(EntityType.TRACT, "human", "yeh2022", local_code)
            definitions.append(
                TractDefinition(
                    id=tract_id,
                    name=f"{full_name} (hemisferio {_HEMISPHERE_LABEL[hemisphere]})",
                    hemisphere=hemisphere,
                    header_code=header_code,
                )
            )
    return definitions


def tract_region_connections(connectome_xlsx_path: Path) -> list[TractRegionConnection]:
    """Lee la matriz tracto-región completa (180 áreas x 26 tractos x 2
    hemisferios = 9360 pares) y devuelve una conexión por cada celda,
    incluidas las de probabilidad 0.0 (sección 24: nunca se descarta un
    dato real al cargarlo). Valida la forma exacta esperada del archivo
    antes de leer ningún valor: nunca se "redondea" sobre un archivo que
    no encaja con lo esperado."""
    import openpyxl

    wb = openpyxl.load_workbook(str(connectome_xlsx_path), read_only=True, data_only=True)
    ws = wb["Tract-Region Connectome"]
    rows = list(ws.iter_rows(min_row=1, max_row=182, values_only=True))
    wb.close()

    if len(rows) != 182:
        raise ValueError(f"se esperaban 182 filas (2 de cabecera + 180 de región), se encontraron {len(rows)}")

    hemisphere_header, tract_header = rows[0], rows[1]
    data_rows = rows[2:182]

    if hemisphere_header[1] != "LEFT":
        raise ValueError(f"se esperaba 'LEFT' en la columna 1, se encontró {hemisphere_header[1]!r}")
    try:
        right_start = hemisphere_header.index("RIGHT")
    except ValueError as exc:
        raise ValueError("no se encontró el bloque 'RIGHT' en la cabecera de hemisferio") from exc

    left_codes = tuple(tract_header[1:right_start])
    right_codes = tuple(tract_header[right_start:])
    if left_codes != EXPECTED_HEADER_TRACT_CODES:
        raise ValueError(f"códigos de tracto (bloque LEFT) inesperados: {left_codes}")
    if right_codes != EXPECTED_HEADER_TRACT_CODES:
        raise ValueError(f"códigos de tracto (bloque RIGHT) inesperados: {right_codes}")

    connections: list[TractRegionConnection] = []
    for row in data_rows:
        area_code = str(row[0]).strip()
        for hemisphere, col_offset in (("L", 1), ("R", right_start)):
            region_local_code = f"{hemisphere}_{area_code}".lower()
            region_id = build_id(EntityType.REGION, "human", "hcp-mmp1", region_local_code)
            for i, header_code in enumerate(EXPECTED_HEADER_TRACT_CODES):
                weight = row[col_offset + i]
                if weight is None:
                    raise ValueError(
                        f"celda vacía inesperada: región {area_code!r}, "
                        f"tracto {header_code!r}, hemisferio {hemisphere!r}"
                    )
                if not (0.0 <= float(weight) <= 1.0):
                    raise ValueError(
                        f"probabilidad fuera de [0,1]: región {area_code!r}, "
                        f"tracto {header_code!r}, hemisferio {hemisphere!r}, valor {weight!r}"
                    )
                tract_local_code = f"{hemisphere}_{header_code}".lower()
                tract_id = build_id(EntityType.TRACT, "human", "yeh2022", tract_local_code)
                conn_id = build_id(
                    EntityType.CONNECTION, "human", "yeh2022",
                    f"{tract_local_code}__{region_local_code}",
                )
                connections.append(
                    TractRegionConnection(
                        id=conn_id,
                        source_id=tract_id,
                        target_id=region_id,
                        weight=float(weight),
                        hemisphere=hemisphere,
                    )
                )
    return connections
